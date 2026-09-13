import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


RUN_DIR = Path(__file__).resolve().parent
SOURCE_TWO_GROUPS = Path(r"D:\DA\待整理\value_added_realted\两个群_20260421-20260801_非标增值讨论_任一艾特命中.xlsx")
THIRD_FILTERED = RUN_DIR / "third_chat_discussions_filtered.json"
THIRD_ALL = RUN_DIR / "third_chat_discussions_all.json"
COMBINED_JSON = RUN_DIR / "three_chat_discussions_combined_filtered.json"
THIRD_FILTERED_NUMBERED_JSON = RUN_DIR / "third_chat_discussions_filtered_numbered.json"
THIRD_ALL_NUMBERED_JSON = RUN_DIR / "third_chat_discussions_all_numbered.json"
SUMMARY_JSON = RUN_DIR / "workbook_summary.json"

COLUMNS = [
    "序号",
    "群名称",
    "群ID",
    "日期",
    "开始时间",
    "结束时间",
    "讨论ID",
    "发起人",
    "增值产品",
    "增值服务",
    "客户/对象",
    "仓库",
    "关联单号",
    "摘要",
    "对话详情",
    "艾特人员",
    "命中艾特人员",
    "参与人",
    "原始消息ID",
    "飞书链接",
    "关键词",
]


def read_source_rows():
    wb = load_workbook(SOURCE_TWO_GROUPS, read_only=True, data_only=True)
    ws = wb["讨论明细"]
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for excel_row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and str(v).strip() for v in excel_row):
            continue
        row = dict(zip(header, excel_row))
        rows.append({col: row.get(col, "") if row.get(col, "") is not None else "" for col in COLUMNS})
    return rows


def renumber(rows, start=1):
    out = []
    for idx, row in enumerate(rows, start):
        item = {col: row.get(col, "") for col in COLUMNS}
        item["序号"] = idx
        out.append(item)
    return out


def count_by(rows, field):
    return Counter(str(row.get(field) or "空/未知") for row in rows).most_common()


def main():
    source_rows = read_source_rows()
    third_filtered = json.loads(THIRD_FILTERED.read_text(encoding="utf-8"))
    third_all = json.loads(THIRD_ALL.read_text(encoding="utf-8"))

    third_filtered_numbered = renumber(third_filtered, 1)
    third_all_numbered = renumber(third_all, 1)
    combined = renumber(source_rows + third_filtered, 1)

    THIRD_FILTERED_NUMBERED_JSON.write_text(
        json.dumps(third_filtered_numbered, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    THIRD_ALL_NUMBERED_JSON.write_text(
        json.dumps(third_all_numbered, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    COMBINED_JSON.write_text(json.dumps(combined, ensure_ascii=False, indent=2), encoding="utf-8")

    summary = {
        "source_two_groups_path": str(SOURCE_TWO_GROUPS),
        "source_two_groups_rows": len(source_rows),
        "third_group_rows_all": len(third_all),
        "third_group_rows_filtered": len(third_filtered),
        "combined_rows_filtered": len(combined),
        "combined_group_distribution": count_by(combined, "群名称"),
        "third_keyword_distribution": count_by(third_filtered, "关键词")[:30],
        "third_rows_with_vasc": sum(1 for r in third_all if "VASC" in str(r.get("关联单号") or "")),
        "third_rows_with_eb": sum(1 for r in third_all if "EB" in str(r.get("关联单号") or "")),
        "third_rows_with_image_marker": sum(1 for r in third_all if "[Image:" in str(r.get("对话详情") or "")),
    }
    SUMMARY_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
